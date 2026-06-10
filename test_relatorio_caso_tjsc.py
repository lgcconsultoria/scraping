#!/usr/bin/env python3
"""Testes offline do relatorio_caso_tjsc (pipeline 4 estágios) contra um stub da API.

Roda com: python test_relatorio_caso_tjsc.py -v

O stub local implementa POST /v1/messages e diferencia os 4 estágios pelo
schema (output_config): perfil tem "teses", extração tem "passagens_verbatim",
aplicação tem "aplicabilidade", síntese não tem output_config. Nenhuma chamada
sai para a internet; nenhuma chave real é usada.
"""

import csv
import importlib
import json
import os
import shutil
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

try:
    import docx as _docx
except ImportError:
    _docx = None

NUM1 = "5000001-11.2025.8.24.0000"
NUM2 = "5000002-22.2025.8.24.0000"
NUM3 = "5000003-33.2025.8.24.0000"

DOC_BASE = ("<html><body>AGRAVO DE INSTRUMENTO. {marca}. " + "lorem juris " * 40
            + "</body></html>")
DOCS = {
    "doc1": DOC_BASE.format(marca="MARCADOR_FAVORAVEL alimentos ex-conjuge"),
    "doc2": DOC_BASE.format(marca="MARCADOR_CONTRARIO sinais exteriores"),
    "doc3": DOC_BASE.format(marca="MARCADOR_IRRELEVANTE alimentos a filho menor"),
}

PERFIL = {
    "teses": [{"id": "t1", "titulo": "Faturamento não é renda",
               "resumo": "...", "precedente_ideal": "..."}],
    "fatos_chave": ["ex-cônjuges sem filhos"], "questoes_juridicas": ["binômio"],
    "termos_busca": ["alimentos"],
}
EXTRACAO_FAV = {"relevante": True, "posicao_preliminar": "favoravel",
                "teses_tocadas": ["t1"], "resultado": "provido ao réu",
                "ratio_decidendi": "faturamento não é renda", "resumo_fatos": "...",
                "passagens_verbatim": ["FATURAMENTO NAO SE CONFUNDE COM RENDA"]}
EXTRACAO_CON = {"relevante": True, "posicao_preliminar": "contrario",
                "teses_tocadas": ["t1"], "resultado": "desprovido",
                "ratio_decidendi": "sinais exteriores", "resumo_fatos": "...",
                "passagens_verbatim": ["PRESUNCAO POR SINAIS EXTERIORES"]}
EXTRACAO_IRR = {"relevante": False, "posicao_preliminar": "neutro",
                "teses_tocadas": [], "resultado": "n/a", "ratio_decidendi": "filhos",
                "resumo_fatos": "...", "passagens_verbatim": []}
APLICACAO_FAV = {"posicao": "favoravel", "aplicabilidade": 9, "teses_apoiadas": ["t1"],
                 "como_se_aplica": "Apoia a tese do binômio.",
                 "trechos_citaveis": [{"trecho": "FATURAMENTO NAO SE CONFUNDE COM RENDA",
                                       "como_usar": "citar"}], "ressalvas": ""}
APLICACAO_CON = {"posicao": "contrario", "aplicabilidade": 3, "teses_apoiadas": ["t1"],
                 "como_se_aplica": "Ameaça a tese.",
                 "trechos_citaveis": [], "ressalvas": "distinguir: havia filhos"}
SINTESE_MD = ("## Sumário executivo\n\nA defesa está **bem amparada** na tese do binômio "
              f"(TJSC, AI n. {NUM1}).\n\n## Precedentes favoráveis\n\n- Forte: caso {NUM1}.")


class StubAnthropic(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_POST(self):
        if self.path.rstrip("/") != "/v1/messages":
            self.send_error(404)
            return
        corpo = json.loads(self.rfile.read(int(self.headers.get("Content-Length", "0"))))
        mensagem = corpo["messages"][0]["content"]
        schema = (corpo.get("output_config") or {}).get("format", {}).get("schema", {})
        props = schema.get("properties", {})
        if "teses" in props and "termos_busca" in props:
            self.server.contagem["perfil"] += 1
            texto = json.dumps(PERFIL)
        elif "passagens_verbatim" in props:
            self.server.contagem["extracao"] += 1
            if "MARCADOR_FAVORAVEL" in mensagem:
                texto = json.dumps(EXTRACAO_FAV)
            elif "MARCADOR_CONTRARIO" in mensagem:
                texto = json.dumps(EXTRACAO_CON)
            else:
                texto = json.dumps(EXTRACAO_IRR)
        elif "aplicabilidade" in props:
            self.server.contagem["aplicacao"] += 1
            texto = json.dumps(APLICACAO_FAV if NUM1 in mensagem else APLICACAO_CON)
        else:
            self.server.contagem["sintese"] += 1
            texto = SINTESE_MD
        resposta = {
            "id": "msg_teste", "type": "message", "role": "assistant",
            "model": corpo.get("model", "claude-opus-4-8"),
            "content": [{"type": "text", "text": texto}],
            "stop_reason": "end_turn", "stop_sequence": None,
            "usage": {"input_tokens": 1000, "output_tokens": 200,
                      "cache_creation_input_tokens": 300, "cache_read_input_tokens": 1500},
        }
        dados = json.dumps(resposta).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(dados)))
        self.end_headers()
        self.wfile.write(dados)


class Servidor(ThreadingHTTPServer):
    daemon_threads = True

    def __init__(self, *args, **kw):
        super().__init__(*args, **kw)
        self.contagem = {"perfil": 0, "extracao": 0, "aplicacao": 0, "sintese": 0}


class TesteRelatorioCaso(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.servidor = Servidor(("127.0.0.1", 0), StubAnthropic)
        porta = cls.servidor.server_address[1]
        threading.Thread(target=cls.servidor.serve_forever, daemon=True).start()
        os.environ["ANTHROPIC_BASE_URL"] = f"http://127.0.0.1:{porta}"
        os.environ["ANTHROPIC_API_KEY"] = "sk-ant-chave-de-teste"
        os.environ["NO_PROXY"] = os.environ["no_proxy"] = "127.0.0.1,localhost"
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        cls.modulo = importlib.import_module("relatorio_caso_tjsc")

    @classmethod
    def tearDownClass(cls):
        cls.servidor.shutdown()
        cls.servidor.server_close()

    def setUp(self):
        self.saida = tempfile.mkdtemp(prefix="tjsc_caso_")
        self.addCleanup(shutil.rmtree, self.saida, ignore_errors=True)
        for chave in self.servidor.contagem:
            self.servidor.contagem[chave] = 0
        os.makedirs(os.path.join(self.saida, "docs"))
        for nome, conteudo in DOCS.items():
            with open(os.path.join(self.saida, "docs", f"{nome}.html"), "w",
                      encoding="utf-8") as arq:
                arq.write(conteudo)
        colunas = ["numero_processo", "relator", "camara", "orgao_julgador", "comarca",
                   "data_julgamento", "classe", "doc_id", "eixo_origem", "formato",
                   "arquivo", "url_integra"]

        def linha(num, doc_id, arquivo):
            return {"numero_processo": num, "relator": "Des. Teste", "camara": "9a_camara",
                    "orgao_julgador": "Nona Câmara", "comarca": "Capital",
                    "data_julgamento": "01/02/2026", "classe": "Agravo de Instrumento",
                    "doc_id": doc_id, "eixo_origem": "binomio_renda", "formato": "html",
                    "arquivo": arquivo, "url_integra": "http://exemplo/x"}

        with open(os.path.join(self.saida, "index.csv"), "w", newline="",
                  encoding="utf-8") as arq:
            escritor = csv.DictWriter(arq, fieldnames=colunas)
            escritor.writeheader()
            escritor.writerows([
                linha(NUM1, "D1" * 15, "docs/doc1.html"),
                linha(NUM2, "D2" * 15, "docs/doc2.html"),
                linha(NUM3, "D3" * 15, "docs/doc3.html"),
                linha("9999999-99.2025.8.24.0000", "D4" * 15, "docs/sumiu.html"),  # arquivo ausente
            ])
        self.caso = os.path.join(self.saida, "caso.md")
        self._escrever_caso("AGRAVO. Tese: faturamento bruto não é renda pessoal. " * 15)

    def _escrever_caso(self, texto):
        with open(self.caso, "w", encoding="utf-8") as arq:
            arq.write(texto)

    def _rodar(self, *flags):
        self.modulo.main(["--caso", self.caso, "--out", self.saida, *flags])

    def _jsonl(self, nome):
        caminho = os.path.join(self.saida, nome)
        if not os.path.exists(caminho):
            return {}
        with open(caminho, encoding="utf-8") as arq:
            return {json.loads(l)["doc_id"]: json.loads(l) for l in arq if l.strip()}

    def test_pipeline_incremental_e_cache(self):
        # 1) execução completa: perfil 1x, extrai os 3 presentes, aplica só os 2 relevantes
        self._rodar()
        self.assertEqual(self.servidor.contagem,
                         {"perfil": 1, "extracao": 3, "aplicacao": 2, "sintese": 1})
        extr = self._jsonl("extracoes_caso.jsonl")
        apli = self._jsonl("aplicacoes_caso.jsonl")
        self.assertEqual(set(extr), {"D1" * 15, "D2" * 15, "D3" * 15})
        self.assertEqual(set(apli), {"D1" * 15, "D2" * 15},
                         "acórdão irrelevante não pode ir para a aplicação (Opus)")
        self.assertEqual(apli["D1" * 15]["aplicacao"]["posicao"], "favoravel")
        self.assertEqual(apli["D2" * 15]["aplicacao"]["posicao"], "contrario")
        self.assertTrue(os.path.exists(os.path.join(self.saida, "perfil_caso.json")))
        with open(os.path.join(self.saida, "relatorio_caso.md"), encoding="utf-8") as arq:
            self.assertIn(NUM1, arq.read())

        # 2) reexecução sem mudanças: NENHUMA chamada à API (tudo em cache)
        self._rodar()
        self.assertEqual(self.servidor.contagem,
                         {"perfil": 1, "extracao": 3, "aplicacao": 2, "sintese": 1})

        # 3) agravo alterado: perfil refeito, extrações REAPROVEITADAS, aplicações refeitas
        self._escrever_caso("AGRAVO v2. Nova redação das teses. " * 20)
        self._rodar()
        self.assertEqual(self.servidor.contagem,
                         {"perfil": 2, "extracao": 3, "aplicacao": 4, "sintese": 2},
                         "mudar o agravo não pode re-extrair acórdãos, só reanalisar")

    def test_acordao_novo_dispara_so_o_novo(self):
        self._rodar()
        base = dict(self.servidor.contagem)
        # acrescenta um acórdão novo ao index.csv (favorável)
        with open(os.path.join(self.saida, "docs", "doc5.html"), "w", encoding="utf-8") as arq:
            arq.write(DOC_BASE.format(marca="MARCADOR_FAVORAVEL novo"))
        caminho_idx = os.path.join(self.saida, "index.csv")
        with open(caminho_idx, encoding="utf-8") as arq:
            conteudo = arq.read()
        with open(caminho_idx, "a", encoding="utf-8") as arq:
            arq.write("5000005-55.2025.8.24.0000,Des. Teste,9a_camara,Nona Câmara,Capital,"
                      "01/02/2026,Agravo de Instrumento,D5D5D5D5D5D5D5D5D5D5D5D5D5D5D5,"
                      "binomio_renda,html,docs/doc5.html,http://exemplo/x\n")
        self._rodar()
        self.assertEqual(self.servidor.contagem["perfil"], base["perfil"], "perfil reusado")
        self.assertEqual(self.servidor.contagem["extracao"], base["extracao"] + 1)
        self.assertEqual(self.servidor.contagem["aplicacao"], base["aplicacao"] + 1)

    def test_limite_extracoes(self):
        self._rodar("--limite", "1")
        self.assertEqual(self.servidor.contagem["extracao"], 1)
        self.assertLessEqual(self.servidor.contagem["aplicacao"], 1)

    @unittest.skipUnless(_docx, "python-docx não instalado")
    def test_gera_docx_formatado(self):
        self._rodar()
        caminho = os.path.join(self.saida, "relatorio_caso.docx")
        self.assertTrue(os.path.exists(caminho))
        documento = _docx.Document(caminho)
        texto = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("Relatório de Precedentes", texto)
        self.assertIn("Sumário executivo", texto)
        self.assertTrue(documento.tables, "o apêndice deve virar uma tabela do Word")
        self.assertIn("Processo", documento.tables[0].rows[0].cells[0].text)

    def test_sem_docx(self):
        self._rodar("--sem-docx")
        self.assertFalse(os.path.exists(os.path.join(self.saida, "relatorio_caso.docx")))
        self.assertTrue(os.path.exists(os.path.join(self.saida, "relatorio_caso.md")))

    def test_exige_chave(self):
        chave = os.environ.pop("ANTHROPIC_API_KEY")
        try:
            with self.assertRaises(SystemExit):
                self._rodar()
        finally:
            os.environ["ANTHROPIC_API_KEY"] = chave


if __name__ == "__main__":
    unittest.main(verbosity=2)
