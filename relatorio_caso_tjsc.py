#!/usr/bin/env python3
"""Relatório de precedentes do caso concreto via API do Claude (pipeline 4 estágios).

Cruza a peça do SEU caso com os acórdãos baixados pelo scraper, em quatro
estágios com modelo e cache pensados para gastar o mínimo de tokens:

  1. PERFIL DO CASO  (Opus, 1x e cacheado em disco)
       Lê o agravo UMA vez e destila um perfil estruturado (teses, fatos-chave,
       questões, o precedente ideal de cada tese). Só roda de novo se o agravo
       mudar (detectado por hash). Os acórdãos NUNCA mais leem o agravo inteiro
       — usam este perfil compacto.

  2. EXTRAÇÃO        (Sonnet — econômico; lê o acórdão COMPLETO)
       Lê o inteiro teor inteiro (sem cortes — o contexto importa) e extrai com
       fidelidade: relevância, ratio decidendi, resultado, fatos e passagens
       VERBATIM citáveis. Roda só em acórdãos novos do index.csv.

  3. APLICAÇÃO       (Opus — raciocínio jurídico; só nos relevantes)
       A partir do perfil + extração, avalia como o acórdão se aplica às teses
       (posição, aplicabilidade 0-10, trechos citáveis, distinguishing). Só os
       acórdãos que a Extração marcou relevantes chegam aqui.

  4. SÍNTESE         (Opus, 1x) -> relatorio_caso.md
       Junta as aplicações num relatório por tese, com parágrafos para a minuta.

Idempotência total: tudo é cacheado em disco. Reexecutar NÃO chama a API se
não houver acórdão novo no index.csv nem alteração no agravo.

Uso:
    export ANTHROPIC_API_KEY="sk-ant-..."        # chave SÓ por variável de ambiente
    python relatorio_caso_tjsc.py --caso agravo.md --limite 5   # teste barato
    python relatorio_caso_tjsc.py --caso agravo.md              # corpus completo
"""

import argparse
import collections
import csv
import datetime
import json
import logging
import os
import re
import sys

from nucleo.cache import (sha256_texto, sha256_arquivo,
                           carregar_jsonl, compactar_jsonl,
                           carregar_json, salvar_json)
from nucleo.custos import Contador, PRECOS_USD

try:
    import anthropic
except ImportError:  # pragma: no cover
    sys.exit("Este script requer o SDK da Anthropic: pip install anthropic")

import triagem_tjsc  # reaproveita a extração de texto (.pdf/.rtf/.html)

OUT_PADRAO = "decisoes_tjsc"
# Opus para raciocínio jurídico (perfil, aplicação, síntese); Sonnet para a
# leitura/extração em massa (mais barato e lê o documento inteiro).
MODELO_PESADO = os.environ.get("TJSC_MODELO_PESADO", "claude-opus-4-8")
MODELO_LEVE = os.environ.get("TJSC_MODELO_LEVE", "claude-sonnet-4-6")

MAX_TOKENS = {"perfil": 8000, "extracao": 8000, "aplicacao": 12000, "sintese": 16000}
MIN_CHARS_TEXTO = 200
TETO_SEGURANCA_CHARS = 1_200_000  # ~360k tokens; acórdão real nunca chega perto
TOP_FAVORAVEIS = 30               # quantas aplicações favoráveis entram na síntese

PERFIL_SCHEMA = {
    "type": "json_schema",
    "schema": {
        "type": "object",
        "properties": {
            "teses": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "string", "description": "identificador curto, ex.: t1"},
                        "titulo": {"type": "string"},
                        "resumo": {"type": "string"},
                        "precedente_ideal": {
                            "type": "string",
                            "description": "que tipo de acórdão apoiaria esta tese.",
                        },
                    },
                    "required": ["id", "titulo", "resumo", "precedente_ideal"],
                    "additionalProperties": False,
                },
            },
            "fatos_chave": {"type": "array", "items": {"type": "string"}},
            "questoes_juridicas": {"type": "array", "items": {"type": "string"}},
            "termos_busca": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["teses", "fatos_chave", "questoes_juridicas", "termos_busca"],
        "additionalProperties": False,
    },
}

EXTRACAO_SCHEMA = {
    "type": "json_schema",
    "schema": {
        "type": "object",
        "properties": {
            "relevante": {
                "type": "boolean",
                "description": "true se a ratio bear em QUALQUER tese do perfil "
                               "(apoiando OU ameaçando); false se for matéria diversa.",
            },
            "posicao_preliminar": {"type": "string",
                                   "enum": ["favoravel", "contrario", "neutro"]},
            "teses_tocadas": {"type": "array", "items": {"type": "string"},
                              "description": "ids das teses do perfil (ex.: t1, t3)."},
            "resultado": {"type": "string",
                          "description": "provido/desprovido/parcial e a quem aproveita."},
            "ratio_decidendi": {"type": "string"},
            "resumo_fatos": {"type": "string"},
            "passagens_verbatim": {
                "type": "array",
                "items": {"type": "string"},
                "description": "Citações LITERAIS do acórdão (cópia exata), candidatas "
                               "a serem citadas na peça. Não parafraseie.",
            },
        },
        "required": ["relevante", "posicao_preliminar", "teses_tocadas", "resultado",
                     "ratio_decidendi", "resumo_fatos", "passagens_verbatim"],
        "additionalProperties": False,
    },
}

APLICACAO_SCHEMA = {
    "type": "json_schema",
    "schema": {
        "type": "object",
        "properties": {
            "posicao": {"type": "string", "enum": ["favoravel", "contrario", "neutro"]},
            "aplicabilidade": {"type": "integer", "enum": list(range(11)),
                               "description": "0-10: identidade fática e jurídica com o caso."},
            "teses_apoiadas": {"type": "array", "items": {"type": "string"}},
            "como_se_aplica": {"type": "string",
                               "description": "raciocínio: como apoia (ou ameaça) cada tese."},
            "trechos_citaveis": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "trecho": {"type": "string",
                                   "description": "cópia LITERAL de uma das passagens_verbatim."},
                        "como_usar": {"type": "string"},
                    },
                    "required": ["trecho", "como_usar"],
                    "additionalProperties": False,
                },
            },
            "ressalvas": {"type": "string",
                          "description": "risco de distinguishing, ou string vazia."},
        },
        "required": ["posicao", "aplicabilidade", "teses_apoiadas", "como_se_aplica",
                     "trechos_citaveis", "ressalvas"],
        "additionalProperties": False,
    },
}

SIST_PERFIL = ("Você é assistente jurídico sênior em direito de família brasileiro e "
               "jurisprudência do TJSC. Analise a peça do Agravante e destile um perfil "
               "estruturado do caso (teses defendidas, fatos-chave, questões jurídicas e "
               "termos de busca) que orientará a avaliação de acórdãos como precedentes. "
               "Seja fiel à peça; não invente teses que ela não sustenta.")

SIST_EXTRACAO = ("Você resume acórdãos do TJSC com FIDELIDADE para um advogado de família. "
                 "Leia o inteiro teor COMPLETO e extraia os elementos pedidos. Uma passagem "
                 "só entra em passagens_verbatim se for cópia LITERAL do texto fornecido — "
                 "nunca parafraseie nem invente. Marque relevante=true se a ratio decidendi "
                 "tocar QUALQUER tese do caso (a favor ou contra); relevante=false para "
                 "matéria diversa (ex.: alimentos a filhos menores, tema sem relação).\n\n"
                 "=== PERFIL DO CASO ===\n{perfil}\n=== FIM DO PERFIL ===")

SIST_APLICACAO = ("Você é assistente jurídico do AGRAVANTE. A partir do perfil do caso e da "
                  "extração de um acórdão (resumo fiel + passagens verbatim), avalie como o "
                  "acórdão se aplica às teses da defesa. REGRAS: (1) 'favoravel' exige que a "
                  "RATIO (não obiter) apoie tese concreta da defesa; um acórdão que nega o "
                  "que o Agravante também pede é 'contrario'. (2) Em trechos_citaveis, copie "
                  "LITERALMENTE apenas passagens presentes na lista passagens_verbatim — não "
                  "crie novas citações. (3) Na dúvida entre favoravel e neutro, escolha neutro "
                  "e explique em ressalvas.\n\n=== PERFIL DO CASO ===\n{perfil}\n=== FIM ===")

SIST_SINTESE = ("Você é assistente jurídico do AGRAVANTE, redigindo um relatório de "
                "precedentes a partir do perfil do caso e das análises (uma por acórdão).\n\n"
                "=== PERFIL DO CASO ===\n{perfil}\n=== FIM DO PERFIL ===")

INSTRUCAO_SINTESE = ("Redija, em markdown e português, um RELATÓRIO DE PRECEDENTES para o "
                     "advogado do Agravante, com: 1) **Sumário executivo** por tese (onde a "
                     "defesa está amparada e onde está exposta); 2) **Precedentes favoráveis, "
                     "por tese**, do mais forte ao mais fraco — cada um com identificação "
                     "completa, por que se aplica, o melhor trecho citável e um PARÁGRAFO "
                     "PRONTO PARA A MINUTA; 3) **Precedentes contrários/de risco** com "
                     "estratégia de distinguishing; 4) **Lacunas** (teses sem respaldo na "
                     "amostra). Cite no formato (TJSC, <classe> n. <número>, rel. <relator>, "
                     "<órgão>, j. <data>). Use somente os acórdãos abaixo; não invente.\n\n"
                     "=== ANÁLISES ===\n{analises}")


# ---------------------------------------------------------------- utilidades

def texto_da_resposta(resposta):
    return "".join(b.text for b in resposta.content if b.type == "text")


def bloco_sistema(texto):
    """Bloco de sistema com prompt caching (o perfil é reusado a ~0,1x do custo)."""
    return [{"type": "text", "text": texto, "cache_control": {"type": "ephemeral"}}]


def chamar(cliente, modelo, sistema, pergunta, contador, schema=None, pensar=True,
           max_tokens=8000):
    kwargs = {
        "model": modelo, "max_tokens": max_tokens, "system": sistema,
        "messages": [{"role": "user", "content": pergunta}],
        "thinking": {"type": "adaptive"} if pensar else {"type": "disabled"},
    }
    if schema:
        kwargs["output_config"] = {"format": schema}
    resposta = cliente.messages.create(**kwargs)
    contador.somar(modelo, resposta.usage)
    return texto_da_resposta(resposta)


# ---------------------------------------------------------------- estágios

def gerar_perfil(cliente, texto_caso, modelo, contador):
    texto = chamar(
        cliente, modelo, bloco_sistema(SIST_PERFIL),
        "Destile o perfil estruturado do caso a partir da peça abaixo.\n\n"
        f"=== PEÇA DO CASO ===\n{texto_caso}\n=== FIM ===",
        contador, schema=PERFIL_SCHEMA, pensar=True, max_tokens=MAX_TOKENS["perfil"])
    return json.loads(texto)


def extrair_acordao(cliente, sistema, linha, texto_acordao, modelo, contador):
    if len(texto_acordao) > TETO_SEGURANCA_CHARS:
        logging.warning("acórdão %s gigante (%s chars); enviando início para caber no contexto",
                        linha["numero_processo"], len(texto_acordao))
        texto_acordao = texto_acordao[:TETO_SEGURANCA_CHARS]
    pergunta = (f"Acórdão — Processo {linha['numero_processo']} | Relator {linha['relator']} "
                f"| {linha['orgao_julgador']} | j. {linha['data_julgamento']} | "
                f"{linha['classe']}.\n\n=== INTEIRO TEOR (completo) ===\n{texto_acordao}")
    # thinking desligado: extração é tarefa direta e queremos economia no modelo leve.
    return json.loads(chamar(cliente, modelo, sistema, pergunta, contador,
                             schema=EXTRACAO_SCHEMA, pensar=False,
                             max_tokens=MAX_TOKENS["extracao"]))


def analisar_aplicacao(cliente, sistema, linha, extracao, modelo, contador):
    pergunta = (f"Acórdão — Processo {linha['numero_processo']} | Relator {linha['relator']} "
                f"| {linha['orgao_julgador']} | j. {linha['data_julgamento']}.\n\n"
                f"=== EXTRAÇÃO (resumo fiel + passagens verbatim) ===\n"
                f"{json.dumps(extracao, ensure_ascii=False, indent=1)}")
    return json.loads(chamar(cliente, modelo, sistema, pergunta, contador,
                             schema=APLICACAO_SCHEMA, pensar=True,
                             max_tokens=MAX_TOKENS["aplicacao"]))


def sintetizar(cliente, sistema, analises, modelo, contador):
    corpo = json.dumps(analises, ensure_ascii=False, indent=1)
    return chamar(cliente, modelo, sistema,
                  INSTRUCAO_SINTESE.format(analises=corpo), contador,
                  pensar=True, max_tokens=MAX_TOKENS["sintese"])


def montar_markdown(texto_sintese, aplicacoes, modelos):
    """Monta o relatório em markdown (síntese da IA + apêndice com a tabela)."""
    ordenadas = sorted(aplicacoes.values(),
                       key=lambda r: (r["aplicacao"]["posicao"] != "favoravel",
                                      -r["aplicacao"]["aplicabilidade"]))
    linhas = [
        "# Relatório de Precedentes — caso concreto × acórdãos TJSC",
        f"> AVISO: relatório gerado por IA ({modelos}) a partir dos acórdãos baixados. "
        "Apoio à decisão — confira cada precedente no inteiro teor antes de citar em juízo.",
        "",
        f"Gerado em: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} | "
        f"Aplicações analisadas: {len(aplicacoes)}",
        "",
        texto_sintese.strip(),
        "",
        "---",
        "",
        "## Apêndice — análises (ordenadas por utilidade)",
        "",
        "| Processo | Posição | Aplic. | Relator | Data | Como se aplica |",
        "|---|---|---|---|---|---|",
    ]
    for reg in ordenadas:
        ap = reg["aplicacao"]
        resumo = ap["como_se_aplica"].replace("|", "/").replace("\n", " ")[:200]
        linhas.append(f"| {reg['numero_processo']} | {ap['posicao']} | "
                      f"{ap['aplicabilidade']}/10 | {reg['relator']} | "
                      f"{reg['data_julgamento']} | {resumo} |")
    return "\n".join(linhas) + "\n"


def escrever_md(caminho, markdown):
    with open(caminho, "w", encoding="utf-8") as arq:
        arq.write(markdown)


_NEGRITO_RE = re.compile(r"\*\*(.+?)\*\*")


def _runs_com_negrito(paragrafo, texto):
    """Adiciona o texto ao parágrafo do Word convertendo **negrito** em runs."""
    pos = 0
    for m in _NEGRITO_RE.finditer(texto):
        if m.start() > pos:
            paragrafo.add_run(texto[pos:m.start()])
        paragrafo.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(texto):
        paragrafo.add_run(texto[pos:])


def escrever_docx(caminho, markdown):
    """Converte o relatório markdown em um .docx formatado (mesma engine da skill
    docx). Sem dependência: se python-docx faltar, retorna False e segue só com .md."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return False

    documento = Document()
    estilo = documento.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    linhas = markdown.splitlines()
    indice = 0
    while indice < len(linhas):
        linha = linhas[indice].rstrip()
        if not linha.strip():
            indice += 1
            continue
        # Tabela markdown: acumula linhas iniciadas por "|"
        if linha.lstrip().startswith("|"):
            bloco = []
            while indice < len(linhas) and linhas[indice].lstrip().startswith("|"):
                bloco.append(linhas[indice].strip())
                indice += 1
            linhas_dados = [l for l in bloco if not re.match(r"^\|[\s:|-]+\|$", l)]
            celulas = [[c.strip() for c in l.strip("|").split("|")] for l in linhas_dados]
            if celulas:
                tabela = documento.add_table(rows=1, cols=len(celulas[0]))
                try:
                    tabela.style = "Light Grid Accent 1"
                except KeyError:
                    tabela.style = "Table Grid"
                for idx, texto in enumerate(celulas[0]):
                    run = tabela.rows[0].cells[idx].paragraphs[0].add_run(texto)
                    run.bold = True
                for linha_dados in celulas[1:]:
                    cels = tabela.add_row().cells
                    for idx, texto in enumerate(linha_dados[:len(celulas[0])]):
                        cels[idx].text = texto
            continue
        if linha.startswith("# "):
            documento.add_heading(linha[2:].strip(), level=0)
        elif linha.startswith("## "):
            documento.add_heading(linha[3:].strip(), level=1)
        elif linha.startswith("### "):
            documento.add_heading(linha[4:].strip(), level=2)
        elif linha.startswith("#### "):
            documento.add_heading(linha[5:].strip(), level=3)
        elif linha.startswith(">"):
            paragrafo = documento.add_paragraph(style="Intense Quote")
            _runs_com_negrito(paragrafo, linha.lstrip("> ").strip())
        elif linha.strip() == "---":
            documento.add_paragraph()
        elif re.match(r"^\s*[-*]\s+", linha):
            paragrafo = documento.add_paragraph(style="List Bullet")
            _runs_com_negrito(paragrafo, re.sub(r"^\s*[-*]\s+", "", linha))
        elif re.match(r"^\s*\d+\.\s+", linha):
            paragrafo = documento.add_paragraph(style="List Number")
            _runs_com_negrito(paragrafo, re.sub(r"^\s*\d+\.\s+", "", linha))
        else:
            _runs_com_negrito(documento.add_paragraph(), linha)
        indice += 1
    documento.save(caminho)
    return True


# ---------------------------------------------------------------- orquestração

def carregar_indice(saida):
    caminho = os.path.join(saida, "index.csv")
    if not os.path.exists(caminho):
        sys.exit(f"{caminho} não encontrado — rode antes o scraper_tjsc.py (com download).")
    with open(caminho, newline="", encoding="utf-8") as arq:
        linhas = [l for l in csv.DictReader(arq) if l.get("doc_id") and l.get("arquivo")]
    caminho_triagem = os.path.join(saida, "triagem.csv")
    if os.path.exists(caminho_triagem):  # prioriza os mais promissores (útil com --limite)
        with open(caminho_triagem, newline="", encoding="utf-8") as arq:
            notas = {}
            for linha in csv.DictReader(arq):
                try:
                    notas[linha["arquivo"]] = max(int(linha["score_eixo1"] or 0),
                                                  int(linha["score_eixo2"] or 0))
                except ValueError:
                    pass
        linhas.sort(key=lambda l: notas.get(l.get("arquivo", ""), -99), reverse=True)
    return linhas


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Pipeline de 4 estágios (perfil/extração/aplicação/síntese) que avalia "
                    "os acórdãos baixados como precedentes para um caso concreto.")
    parser.add_argument("--caso", required=True,
                        help="peça do caso (md/txt), fora do repositório")
    parser.add_argument("--out", default=OUT_PADRAO, help=f"pasta do scraper (padrão: {OUT_PADRAO})")
    parser.add_argument("--limite", type=int, default=0,
                        help="extrai no máximo N acórdãos NOVOS nesta execução (0 = todos)")
    parser.add_argument("--modelo-pesado", default=MODELO_PESADO,
                        help=f"modelo de raciocínio (padrão: {MODELO_PESADO})")
    parser.add_argument("--modelo-leve", default=MODELO_LEVE,
                        help=f"modelo de extração (padrão: {MODELO_LEVE})")
    parser.add_argument("--forcar-relatorio", action="store_true",
                        help="regenera o relatório a partir do cache mesmo sem novidades")
    parser.add_argument("--sem-docx", action="store_true",
                        help="gera apenas o .md, sem o .docx formatado (requer python-docx)")
    parser.add_argument("--refazer", action="store_true",
                        help="ignora TODO o cache e refaz perfil, extrações e análises (re-cobra!)")
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", force=True)
    if not (os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")):
        sys.exit('Defina ANTHROPIC_API_KEY antes de rodar (export ANTHROPIC_API_KEY="sk-ant-...").')

    with open(args.caso, encoding="utf-8") as arq:
        texto_caso = arq.read()
    hash_agravo = sha256_texto(texto_caso)

    p_perfil = os.path.join(args.out, "perfil_caso.json")
    p_extr = os.path.join(args.out, "extracoes_caso.jsonl")
    p_apli = os.path.join(args.out, "aplicacoes_caso.jsonl")
    p_estado = os.path.join(args.out, "relatorio_estado.json")
    p_md = os.path.join(args.out, "relatorio_caso.md")
    p_docx = os.path.join(args.out, "relatorio_caso.docx")
    if args.refazer:
        for caminho in (p_perfil, p_extr, p_apli, p_estado):
            if os.path.exists(caminho):
                os.remove(caminho)

    # --- estado em disco (sem tocar na API ainda) ---
    perfil_meta = carregar_json(p_perfil)
    precisa_perfil = (perfil_meta is None or perfil_meta.get("hash_agravo") != hash_agravo)
    extracoes = carregar_jsonl(p_extr)
    aplicacoes_todas = carregar_jsonl(p_apli)
    # aplicações só valem para o agravo atual (se o agravo mudou, são refeitas)
    aplicacoes = {k: v for k, v in aplicacoes_todas.items()
                  if v.get("hash_agravo") == hash_agravo}

    documentos = carregar_indice(args.out)
    hashes = {}  # doc_id -> hash do arquivo atual (re-extrai se o arquivo mudou)
    presentes = []
    for linha in documentos:
        caminho = os.path.join(args.out, linha["arquivo"])
        if os.path.exists(caminho):
            hashes[linha["doc_id"]] = sha256_arquivo(caminho)
            presentes.append(linha)

    def extracao_valida(doc_id):
        reg = extracoes.get(doc_id)
        return reg is not None and reg.get("hash_arquivo") == hashes.get(doc_id)

    def eh_relevante(doc_id):
        reg = extracoes.get(doc_id)
        return bool(reg and reg["extracao"].get("relevante"))

    pend_extracao = [l for l in presentes if not extracao_valida(l["doc_id"])]
    relevantes = [l for l in presentes if extracao_valida(l["doc_id"]) and eh_relevante(l["doc_id"])]
    pend_aplicacao = [l for l in relevantes if l["doc_id"] not in aplicacoes]
    ids_relevantes = sorted(l["doc_id"] for l in relevantes)
    estado = carregar_json(p_estado)
    relatorio_ok = (os.path.exists(p_md) and estado is not None
                    and estado.get("hash_agravo") == hash_agravo
                    and estado.get("ids_relevantes") == ids_relevantes)

    if (not precisa_perfil and not pend_extracao and not pend_aplicacao
            and relatorio_ok and not args.forcar_relatorio):
        logging.info("Nada novo: agravo inalterado e nenhum acórdão novo no index.csv.")
        print("Nada a fazer — o relatório já está atualizado (nenhuma chamada à API). "
              f"Veja {p_md}")
        return

    cliente = anthropic.Anthropic(max_retries=4)
    contador = Contador()
    estagios = collections.Counter()
    if args.limite > 0:
        pend_extracao = pend_extracao[:args.limite]
    logging.info("agravo %s | perfil %s | extrações novas: %s | aplicações pendentes: %s",
                 "alterado/novo" if precisa_perfil else "em cache",
                 "regerar" if precisa_perfil else "ok", len(pend_extracao), len(pend_aplicacao))

    try:
        # 1) PERFIL (Opus) — uma vez; reusa do disco se o agravo não mudou
        if precisa_perfil:
            logging.info("gerando perfil do caso (%s)...", args.modelo_pesado)
            perfil = gerar_perfil(cliente, texto_caso, args.modelo_pesado, contador)
            salvar_json(p_perfil, {"hash_agravo": hash_agravo, "modelo": args.modelo_pesado,
                                   "gerado_em": datetime.datetime.now().isoformat(timespec="seconds"),
                                   "perfil": perfil})
            estagios["perfil"] += 1
            aplicacoes = {}  # agravo mudou -> análises antigas não valem mais
        else:
            perfil = perfil_meta["perfil"]
        logging.info("Perfil: %s teses — %s",
                     len(perfil.get("teses", [])),
                     " | ".join(f"{t['id']}: {t['titulo']}"
                                for t in perfil.get("teses", [])))
        perfil_str = json.dumps(perfil, ensure_ascii=False, indent=1)
        sist_extr = bloco_sistema(SIST_EXTRACAO.format(perfil=perfil_str))
        sist_apli = bloco_sistema(SIST_APLICACAO.format(perfil=perfil_str))

        # 2) EXTRAÇÃO (Sonnet) — lê o acórdão COMPLETO, só os novos/alterados
        for indice, linha in enumerate(pend_extracao, 1):
            rotulo = linha["numero_processo"] or linha["doc_id"]
            try:
                texto = triagem_tjsc.extrair_texto(os.path.join(args.out, linha["arquivo"]))
                if len(texto.strip()) < MIN_CHARS_TEXTO:
                    logging.warning("[extração %s/%s] %s: sem texto extraível, pulando",
                                    indice, len(pend_extracao), rotulo)
                    continue
                logging.info("[extração %s/%s] %s: %s chars → enviando ao modelo",
                             indice, len(pend_extracao), rotulo, len(texto))
                extracao = extrair_acordao(cliente, sist_extr, linha, texto,
                                           args.modelo_leve, contador)
                registro = {chave: linha.get(chave, "") for chave in
                            ("doc_id", "numero_processo", "relator", "camara",
                             "orgao_julgador", "data_julgamento", "classe", "arquivo")}
                registro.update(hash_arquivo=hashes[linha["doc_id"]], extracao=extracao)
                extracoes[linha["doc_id"]] = registro
                estagios["extracao"] += 1
                teses = extracao.get("teses_tocadas") or []
                ratio = (extracao.get("ratio_decidendi") or "").replace("\n", " ")[:200]
                logging.info("[extração %s/%s] %s -> relevante=%s (%s) teses=%s | ~US$ %.2f",
                             indice, len(pend_extracao), rotulo, extracao["relevante"],
                             extracao["posicao_preliminar"], teses or "nenhuma", contador.custo())
                if ratio:
                    logging.info("  ratio: %s", ratio)
            except Exception as exc:
                estagios["falhas"] += 1
                logging.error("[extração] %s: %s", rotulo, exc)
        compactar_jsonl(p_extr, extracoes)

        # recomputa relevantes/pendentes após as novas extrações
        relevantes = [l for l in presentes
                      if extracao_valida(l["doc_id"]) and eh_relevante(l["doc_id"])]
        pend_aplicacao = [l for l in relevantes if l["doc_id"] not in aplicacoes]

        # 3) APLICAÇÃO (Opus) — só nos relevantes ainda sem análise para este agravo
        for indice, linha in enumerate(pend_aplicacao, 1):
            rotulo = linha["numero_processo"] or linha["doc_id"]
            try:
                extracao = extracoes[linha["doc_id"]]["extracao"]
                aplicacao = analisar_aplicacao(cliente, sist_apli, linha, extracao,
                                               args.modelo_pesado, contador)
                registro = {chave: linha.get(chave, "") for chave in
                            ("doc_id", "numero_processo", "relator", "camara",
                             "orgao_julgador", "data_julgamento", "classe", "arquivo")}
                registro.update(hash_agravo=hash_agravo, aplicacao=aplicacao)
                aplicacoes[linha["doc_id"]] = registro
                estagios["aplicacao"] += 1
                logging.info("[aplicação %s/%s] %s -> %s (%s/10) | ~US$ %.2f",
                             indice, len(pend_aplicacao), rotulo, aplicacao["posicao"],
                             aplicacao["aplicabilidade"], contador.custo())
            except Exception as exc:
                estagios["falhas"] += 1
                logging.error("[aplicação] %s: %s", rotulo, exc)
        compactar_jsonl(p_apli, aplicacoes)
    except anthropic.AuthenticationError:
        sys.exit("Chave de API inválida/revogada (401). Gere uma nova e exporte ANTHROPIC_API_KEY.")
    except KeyboardInterrupt:
        compactar_jsonl(p_extr, extracoes)
        compactar_jsonl(p_apli, aplicacoes)
        logging.warning("interrompido; progresso salvo. Reexecute para continuar.")
        sys.exit(130)

    # 4) SÍNTESE (Opus) — só se houve novidade ou o relatório está desatualizado
    ids_relevantes = sorted(l["doc_id"] for l in relevantes)
    precisa_sintese = (estagios["aplicacao"] > 0 or precisa_perfil or not os.path.exists(p_md)
                       or args.forcar_relatorio
                       or (estado or {}).get("ids_relevantes") != ids_relevantes)
    if precisa_sintese and aplicacoes:
        sist_sint = bloco_sistema(SIST_SINTESE.format(
            perfil=json.dumps(perfil, ensure_ascii=False, indent=1)))
        favoraveis = sorted((r for r in aplicacoes.values()
                             if r["aplicacao"]["posicao"] == "favoravel"),
                            key=lambda r: -r["aplicacao"]["aplicabilidade"])[:TOP_FAVORAVEIS]
        contrarios = [r for r in aplicacoes.values()
                      if r["aplicacao"]["posicao"] == "contrario"]
        logging.info("sintetizando relatório (%s favoráveis, %s contrários)...",
                     len(favoraveis), len(contrarios))
        texto_sintese = sintetizar(cliente, sist_sint,
                                   {"favoraveis": favoraveis, "contrarios_ou_risco": contrarios},
                                   args.modelo_pesado, contador)
        markdown = montar_markdown(texto_sintese, aplicacoes,
                                   f"{args.modelo_pesado} + {args.modelo_leve}")
        escrever_md(p_md, markdown)
        if not args.sem_docx and not escrever_docx(p_docx, markdown):
            logging.warning("python-docx ausente — gerei só o .md (pip install python-docx "
                            "para o .docx formatado).")
        salvar_json(p_estado, {"hash_agravo": hash_agravo, "ids_relevantes": ids_relevantes,
                               "gerado_em": datetime.datetime.now().isoformat(timespec="seconds")})
        estagios["sintese"] += 1
    elif not aplicacoes:
        logging.warning("nenhuma aplicação disponível — relatório não gerado.")

    favs = sum(1 for r in aplicacoes.values() if r["aplicacao"]["posicao"] == "favoravel")
    contras = sum(1 for r in aplicacoes.values() if r["aplicacao"]["posicao"] == "contrario")
    print(f"\nConcluído. perfil={estagios['perfil']} | extrações novas={estagios['extracao']} "
          f"| aplicações novas={estagios['aplicacao']} | sínteses={estagios['sintese']} "
          f"| falhas={estagios['falhas']}")
    print(f"Acervo analisado: {len(aplicacoes)} aplicações ({favs} favoráveis, {contras} contrários) "
          f"| relevantes={len(relevantes)} de {len(presentes)} acórdãos")
    detalhe = " | ".join(
        f"{m}: in={u['entrada']:,} cw={u['cache_escrita']:,} cr={u['cache_leitura']:,} out={u['saida']:,}"
        for m, u in contador.por_modelo.items())
    print(f"Tokens [{detalhe or 'nenhuma chamada'}] | custo estimado ~US$ {contador.custo():.2f}")
    print(f"Relatório: {p_md}" + (f" e {p_docx}" if os.path.exists(p_docx) else ""))


if __name__ == "__main__":
    main()
